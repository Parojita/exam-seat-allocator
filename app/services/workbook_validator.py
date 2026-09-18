from openpyxl import load_workbook


STUDENT_REQUIRED_COLUMNS = (
    "roll_number",
    "name",
    "branch_code",
    "semester",
    "email",
    "photo_path",
    "special_request",
    "rfid_uid",
)

STUDENT_REQUIRED_VALUES = (
    "roll_number",
    "name",
    "branch_code",
    "semester",
    "email",
)


def is_blank(value):
    return value is None or str(value).strip() == ""


def validate_semester(value):
    try:
        semester = int(value)
    except (TypeError, ValueError):
        return False

    return 1 <= semester <= 8


def validate_student_workbook(workbook_file):
    workbook = load_workbook(
        workbook_file,
        read_only=True,
        data_only=True,
    )

    worksheet = workbook.active

    first_row = next(
        worksheet.iter_rows(
            min_row=1,
            max_row=1,
            values_only=True,
        ),
        (),
    )

    column_names = [
        str(value).strip()
        if value is not None
        else ""
        for value in first_row
    ]

    columns = {
        column
        for column in column_names
        if column
    }

    missing_columns = [
        column
        for column in STUDENT_REQUIRED_COLUMNS
        if column not in columns
    ]

    column_positions = {
        column: position
        for position, column in enumerate(column_names)
        if column
    }

    row_count = 0
    row_errors = []
    observed_roll_numbers = set()

    for row_number, row in enumerate(
        worksheet.iter_rows(
            min_row=2,
            values_only=True,
        ),
        start=2,
    ):
        contains_data = any(
            not is_blank(value)
            for value in row
        )

        if not contains_data:
            continue

        row_count += 1

        if missing_columns:
            continue

        messages = []

        row_values = {
            column: (
                row[position]
                if position < len(row)
                else None
            )
            for column, position in column_positions.items()
        }

        for column in STUDENT_REQUIRED_VALUES:
            if is_blank(row_values.get(column)):
                messages.append(
                    f"{column} is required."
                )

        semester = row_values.get("semester")

        if (
            not is_blank(semester)
            and not validate_semester(semester)
        ):
            messages.append(
                "semester must be between 1 and 8."
            )

        roll_number = row_values.get("roll_number")

        if not is_blank(roll_number):
            normalized_roll_number = (
                str(roll_number).strip().casefold()
            )

            if normalized_roll_number in observed_roll_numbers:
                messages.append(
                    "roll_number is duplicated in the workbook."
                )
            else:
                observed_roll_numbers.add(
                    normalized_roll_number
                )

        if messages:
            row_errors.append(
                {
                    "row_number": row_number,
                    "messages": messages,
                }
            )

    workbook.close()

    return {
        "valid": (
            not missing_columns
            and not row_errors
        ),
        "row_count": row_count,
        "missing_columns": missing_columns,
        "row_errors": row_errors,
    }