from openpyxl import load_workbook

from app.extensions import db
from app.models import Branch, Student


def clean_optional_value(value):
    if value is None:
        return None

    cleaned_value = str(value).strip()

    if cleaned_value == "":
        return None

    return cleaned_value


def import_student_workbook(workbook_file):
    workbook = load_workbook(
        workbook_file,
        read_only=True,
        data_only=True,
    )

    try:
        worksheet = workbook.active

        first_row = next(
            worksheet.iter_rows(
                min_row=1,
                max_row=1,
                values_only=True,
            )
        )

        column_positions = {
            str(value).strip(): position
            for position, value in enumerate(first_row)
            if value is not None
        }

        existing_roll_numbers = {
            roll_number.strip().casefold()
            for roll_number in db.session.execute(
                db.select(Student.roll_number)
            ).scalars()
        }

        branches = {
            branch.code.strip().casefold(): branch
            for branch in db.session.execute(
                db.select(Branch)
            ).scalars()
        }

        imported_count = 0
        skipped_count = 0

        for row in worksheet.iter_rows(
            min_row=2,
            values_only=True,
        ):
            contains_data = any(
                value is not None
                and str(value).strip() != ""
                for value in row
            )

            if not contains_data:
                continue

            values = {
                column: (
                    row[position]
                    if position < len(row)
                    else None
                )
                for column, position
                in column_positions.items()
            }

            roll_number = str(
                values["roll_number"]
            ).strip()

            normalized_roll_number = (
                roll_number.casefold()
            )

            if normalized_roll_number in existing_roll_numbers:
                skipped_count += 1
                continue

            branch_code = str(
                values["branch_code"]
            ).strip()

            normalized_branch_code = (
                branch_code.casefold()
            )

            branch = branches.get(
                normalized_branch_code
            )

            if branch is None:
                branch = Branch(
                    code=branch_code,
                    name=branch_code,
                )

                db.session.add(branch)

                branches[
                    normalized_branch_code
                ] = branch

            student = Student(
                roll_number=roll_number,
                name=str(values["name"]).strip(),
                branch=branch,
                semester=int(values["semester"]),
                email=clean_optional_value(
                    values.get("email")
                ),
                photo_path=clean_optional_value(
                    values.get("photo_path")
                ),
                special_request=clean_optional_value(
                    values.get("special_request")
                ),
                rfid_uid=clean_optional_value(
                    values.get("rfid_uid")
                ),
            )

            db.session.add(student)

            existing_roll_numbers.add(
                normalized_roll_number
            )

            imported_count += 1

        db.session.commit()

        return {
            "imported_count": imported_count,
            "skipped_count": skipped_count,
        }

    except Exception:
        db.session.rollback()
        raise

    finally:
        workbook.close()